"""Generate the editable EdgeAIKWS SIH jury presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "EdgeAIKWS_SIH26172.pptx"
EMU_PER_INCH = 914400
WIDTH_IN = 20.0
HEIGHT_IN = 11.25

NAVY = "1F3864"
INK = "262F38"
SLATE = "596979"
MUTED = "778491"
ACCENT = "B85D32"
PALE = "F2F4F6"
PALE_BLUE = "EAF0F5"
LINE = "D6DDE3"
WHITE = "FFFFFF"
FONT = "Arial"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, line_width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line) if line else rgb(fill)
    shape.line.width = Pt(line_width)
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    value,
    size=16,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    line_spacing=1.05,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(max(2, size * 0.22))
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.4, dashed=False, arrow=False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        end = parse_xml(
            f'<a:tailEnd {nsdecls("a")} type="triangle" w="sm" len="sm"/>'
        )
        line._element.spPr.ln.append(end)
    return line


def add_card(slide, x, y, w, h, title, body, accent=False, fill=WHITE,
             title_size=17, body_size=13, title_color=NAVY):
    add_rect(slide, x, y, w, h, fill, LINE, 1.0)
    if accent:
        add_rect(slide, x, y, 0.07, h, ACCENT, ACCENT, 0)
    left = x + (0.18 if accent else 0.14)
    add_text(slide, left, y + 0.12, w - (left - x) - 0.12, 0.5,
             title, title_size, title_color, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, left, y + 0.68, w - (left - x) - 0.15, h - 0.79,
             body, body_size, SLATE)


def add_header(slide, title, section, number):
    add_text(slide, 0.78, 0.35, 10, 0.25, section.upper(), 10, ACCENT, True)
    add_text(slide, 0.78, 0.62, 18.4, 0.58, title, 29, NAVY, True)
    add_rect(slide, 0.8, 1.3, 18.4, 0.018, LINE, LINE, 0)
    add_text(slide, 0.8, 10.83, 13.5, 0.2,
             "EDGEAIKWS  /  SIH 26172  /  ISRO · DEPARTMENT OF SPACE",
             9, MUTED)
    add_text(slide, 18.5, 10.83, 0.7, 0.2, f"{number:02d} / 08",
             9, MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    return slide


def add_step(slide, x, y, w, n, title, body, active=False):
    add_rect(slide, x, y, w, 1.9, PALE_BLUE if active else PALE,
             ACCENT if active else LINE, 1.25)
    add_text(slide, x + 0.16, y + 0.12, 0.42, 0.3, f"{n:02d}",
             12, ACCENT if active else MUTED, True)
    add_text(slide, x + 0.16, y + 0.51, w - 0.32, 0.43, title,
             16, NAVY, True)
    add_text(slide, x + 0.16, y + 0.98, w - 0.32, 0.72, body,
             12.5, SLATE)


def slide_1(prs):
    slide = new_slide(prs)
    add_text(slide, 0.92, 0.72, 11.2, 0.3,
             "SMART INDIA HACKATHON 2026  ·  ISRO  ·  HARDWARE",
             12, ACCENT, True)
    add_text(slide, 0.9, 1.55, 12.2, 1.05, "EdgeAIKWS", 50, NAVY, True)
    add_text(slide, 0.95, 2.72, 11.7, 0.78,
             "On-device wake-word detection\nwith selective speech streaming",
             26, INK, False, line_spacing=1.0)
    add_rect(slide, 0.95, 4.08, 10.7, 0.025, ACCENT, ACCENT, 0)
    add_text(slide, 0.95, 4.38, 10.7, 1.05,
             "Low Latency and Efficient Voice Activator for Edge Devices",
             21, NAVY, True)
    add_text(slide, 0.97, 5.55, 3.8, 0.34, "PROBLEM STATEMENT", 10, MUTED, True)
    add_text(slide, 0.95, 5.91, 3.7, 0.72, "26172", 32, NAVY, True)

    metadata = [
        (0.95, "ORGANIZATION", "ISRO / Department of Space"),
        (4.65, "CATEGORY", "Hardware"),
        (7.65, "THEME", "Smart Automation"),
    ]
    for x, label, value in metadata:
        add_text(slide, x, 7.2, 3.4, 0.25, label, 10, MUTED, True)
        add_text(slide, x, 7.55, 3.45, 0.52, value, 16, INK, True)

    # Compact, editable system path as the cover's only diagram.
    add_rect(slide, 13.15, 1.55, 5.8, 6.15, PALE, LINE, 1.0)
    add_text(slide, 13.55, 1.9, 4.9, 0.35, "EDGE-TO-CLOUD FLOW", 11, MUTED, True)
    flow = [
        (13.55, 2.58, "INMP441", "Microphone"),
        (13.55, 3.72, "ESP32-S3", "Local Hello Tors KWS"),
        (13.55, 4.86, "WebSocket", "Triggered PCM stream"),
        (13.55, 6.0, "Backend", "ASR processing"),
    ]
    for index, (x, y, title, body) in enumerate(flow):
        add_rect(slide, x, y, 4.95, 0.78, WHITE,
                 ACCENT if index == 2 else LINE, 1.2)
        add_text(slide, x + 0.16, y + 0.08, 1.7, 0.28,
                 title, 15, NAVY, True)
        add_text(slide, x + 1.92, y + 0.1, 2.8, 0.52,
                 body, 13, SLATE, valign=MSO_ANCHOR.MIDDLE)
        if index < len(flow) - 1:
            add_line(slide, 16.03, y + 0.79, 16.03, y + 1.1,
                     ACCENT if index == 1 else MUTED, 1.5, arrow=True)
    add_text(slide, 13.55, 7.02, 4.9, 0.42,
             "Audio upload begins only after local detection.",
             12, SLATE)
    add_text(slide, 0.95, 10.83, 14.0, 0.2,
             "CUSTOM KEYWORD: HELLO TORS  ·  ESP32-S3  ·  TFLITE MICRO",
             9, MUTED, True)
    add_text(slide, 18.5, 10.83, 0.7, 0.2, "01 / 08", 9, MUTED,
             align=PP_ALIGN.RIGHT)
    return slide


def slide_2(prs):
    slide = new_slide(prs)
    add_header(slide, "Voice activation should start at the edge",
               "01  /  Problem and requirements", 2)
    add_text(slide, 0.9, 1.65, 8.2, 1.25,
             "Always streaming speech to the cloud increases data exposure, network use and response delay.\n\n"
             "The required design listens locally for a custom phrase, then streams relevant speech to remote ASR.",
             19, INK, line_spacing=1.12)
    add_text(slide, 0.92, 4.45, 7.8, 0.3, "SIH DESIGN CONSTRAINTS", 11, MUTED, True)
    constraints = [
        (0.92, 4.95, "OPEN SOURCE", "TinyML framework; no commercial voice-activation SDK"),
        (4.8, 4.95, "CUSTOM KEYWORD", "No generic assistant trigger; phrase: Hello Tors"),
    ]
    for x, y, title, body in constraints:
        add_card(slide, x, y, 3.65, 1.55, title, body,
                 title_size=14, body_size=12)

    add_text(slide, 9.85, 1.66, 8.9, 0.32, "EVALUATION TARGETS", 11, MUTED, True)
    rows = [
        ("EFFICIENCY", "<256 KB RAM", "<10% idle CPU"),
        ("ACCURACY", "High true-positive rate", "Near-zero false activations"),
        ("LATENCY", "Keyword end → server receives audio", "Report p50 / p95 / p99"),
    ]
    y = 2.15
    for i, (label, value, detail) in enumerate(rows):
        add_rect(slide, 9.82, y, 8.85, 1.18,
                 PALE if i % 2 == 0 else WHITE, LINE, 0.9)
        add_text(slide, 10.08, y + 0.15, 1.8, 0.28, label,
                 11, ACCENT, True)
        add_text(slide, 12.05, y + 0.11, 3.05, 0.36,
                 value, 15, NAVY, True)
        add_text(slide, 15.15, y + 0.12, 3.25, 0.52,
                 detail, 12.5, SLATE, valign=MSO_ANCHOR.MIDDLE)
        y += 1.35
    add_rect(slide, 0.92, 7.12, 17.75, 1.36, PALE_BLUE, LINE, 0.9)
    add_text(slide, 1.18, 7.37, 2.1, 0.32, "EDGE-FIRST", 12, NAVY, True)
    add_text(slide, 3.16, 7.29, 15.0, 0.75,
             "The ESP32 decides locally whether a wake event occurred. The server receives audio only after that event;"
             " RAM, CPU, false-activation rate and end-to-end latency still require measured validation.",
             14, INK, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.95, 9.08, 17.6, 0.5,
             "Target limits are SIH requirements—not results claimed for the current prototype.",
             12, MUTED)
    return slide


def slide_3(prs):
    slide = new_slide(prs)
    add_header(slide, "Local wake-up; conditional audio upload",
               "02  /  System architecture", 3)
    # Trust/compute boundaries.
    add_rect(slide, 0.82, 1.62, 13.0, 6.45, PALE, LINE, 1.0)
    add_rect(slide, 14.2, 1.62, 4.98, 6.45, PALE_BLUE, LINE, 1.0)
    add_text(slide, 1.12, 1.83, 8.5, 0.32,
             "EDGE  ·  ESP32-S3 N16R8  ·  LOCAL INFERENCE", 11, NAVY, True)
    add_text(slide, 14.5, 1.83, 4.2, 0.32,
             "BACKEND  ·  REMOTE ASR", 11, NAVY, True)

    edge = [
        (1.13, "INMP441", "I²S mic\n16 kHz mono"),
        (3.65, "PCM capture", "PCM16\nring buffer"),
        (6.17, "MFCC front end", "49 × 40\nINT8 features"),
        (8.69, "Hello Tors KWS", "TFLite Micro\nlocal decision"),
        (11.21, "Audio stream", "Prebuffer + live\nafter trigger"),
    ]
    for idx, (x, title, body) in enumerate(edge):
        active = idx in (3, 4)
        add_rect(slide, x, 3.15, 2.17, 1.55,
                 PALE_BLUE if active else WHITE,
                 ACCENT if idx == 4 else LINE, 1.2)
        add_text(slide, x + 0.12, 3.42, 1.94, 0.45,
                 title, 14, NAVY, True, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.1, 3.92, 1.97, 0.58,
                 body, 12, SLATE, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
        if idx < len(edge) - 1:
            add_line(slide, x + 2.18, 3.92, x + 2.45, 3.92,
                     ACCENT if idx >= 2 else MUTED, 1.5, arrow=True)
    # Network crossing and server-side ASR.
    add_line(slide, 13.4, 3.92, 14.45, 3.92, ACCENT, 2.0, arrow=True)
    add_text(slide, 14.58, 2.45, 4.15, 0.34,
             "WI-FI / WEBSOCKET", 11, ACCENT, True, align=PP_ALIGN.CENTER)
    add_card(slide, 14.58, 3.15, 4.15, 1.55,
             "PCM receiver", "Validate header, sequence and stream; reconstruct 16-kHz PCM.",
             title_size=15, body_size=12.5, fill=WHITE)
    add_line(slide, 16.65, 4.75, 16.65, 5.08, MUTED, 1.4, arrow=True)
    add_card(slide, 14.58, 5.12, 4.15, 1.47,
             "ASR worker", "Separate backend source includes Vosk ASR; results are logged server-side.",
             title_size=15, body_size=12.3, fill=WHITE)
    add_text(slide, 1.15, 5.2, 11.8, 0.62,
             "Idle microphone audio stays local. A positive KWS decision opens the upload session;"
             " the prebuffer can include part of the wake phrase.",
             14, INK, valign=MSO_ANCHOR.MIDDLE)
    # Separate local playback test from the ASR path; it is not server-generated speech.
    add_text(slide, 0.98, 8.38, 17.5, 0.24,
             "SEPARATE ASR-FREE PLAYBACK TEST  ·  NOT PART OF THE ASR PIPELINE",
             10, MUTED, True)
    add_rect(slide, 0.98, 8.7, 5.25, 0.75, WHITE, LINE, 0.9)
    add_text(slide, 1.12, 8.76, 4.95, 0.25,
             "Local test sender", 12.5, NAVY, True)
    add_text(slide, 1.12, 9.07, 4.95, 0.25,
             "local_backend.py · --play-wav / --test-tone", 10.5, SLATE)
    add_line(slide, 6.3, 9.08, 7.35, 9.08, ACCENT, 1.4, arrow=True)
    add_rect(slide, 7.48, 8.7, 3.65, 0.75, WHITE, LINE, 0.9)
    add_text(slide, 7.62, 8.76, 3.35, 0.25,
             "WebSocket PCM", 12.5, NAVY, True)
    add_text(slide, 7.62, 9.07, 3.35, 0.25,
             "16 kHz mono frames", 10.5, SLATE)
    add_line(slide, 11.2, 9.08, 12.25, 9.08, ACCENT, 1.4, arrow=True)
    add_rect(slide, 12.38, 8.7, 5.95, 0.75, WHITE, LINE, 0.9)
    add_text(slide, 12.52, 8.76, 5.65, 0.25,
             "MAX98357A / speaker", 12.5, NAVY, True)
    add_text(slide, 12.52, 9.07, 5.65, 0.25,
             "Firmware path exists · acoustic output pending", 10.5, SLATE)
    return slide


def slide_4(prs):
    slide = new_slide(prs)
    add_header(slide, "Custom keyword model, deployed without changing its weights",
               "03  /  Model and inference", 4)
    add_text(slide, 0.93, 1.58, 5.1, 0.3,
             "MODEL ARTIFACT  ·  PREPARED BY ML TEAM", 10.5, MUTED, True)
    add_card(slide, 0.92, 2.02, 4.8, 1.25,
             "Custom phrase", "Hello Tors\nNo generic assistant keyword", accent=True,
             title_size=16, body_size=12.5)
    add_line(slide, 5.82, 2.64, 6.25, 2.64, ACCENT, 1.5, arrow=True)
    add_card(slide, 6.34, 2.02, 4.8, 1.25,
             "INT8 TFLite model", "36,744 bytes\nModel artifact embedded unchanged",
             title_size=16, body_size=12.5)
    add_line(slide, 11.25, 2.64, 11.68, 2.64, ACCENT, 1.5, arrow=True)
    add_card(slide, 11.77, 2.02, 6.95, 1.25,
             "ESP32-S3 runtime", "Open-source ESP-IDF + TensorFlow Lite Micro",
             title_size=16, body_size=13)

    add_text(slide, 0.93, 3.75, 5.5, 0.3,
             "LIVE INFERENCE PATH  ·  ESP32-S3", 10.5, MUTED, True)
    flow = [
        (0.93, "PCM16", "16 kHz mono\nnormalize by 32768"),
        (4.55, "Frame", "480 samples\n30 ms; 320 hop"),
        (8.17, "Spectrum", "Periodic Hann\n512-point FFT magnitude"),
        (11.79, "MFCC", "40 HTK mel bands\nlog-mel + exported DCT"),
        (15.41, "Quantized input", "49 × 40 INT8\n≈ 990 ms context"),
    ]
    for i, (x, title, body) in enumerate(flow):
        add_rect(slide, x, 4.35, 3.18, 1.55,
                 PALE_BLUE if i == 4 else WHITE,
                 ACCENT if i == 4 else LINE, 1.1)
        add_text(slide, x + 0.13, 4.62, 2.92, 0.35,
                 title, 15, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.12, 5.06, 2.94, 0.63,
                 body, 12.5, SLATE, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
        if i < len(flow) - 1:
            add_line(slide, x + 3.19, 5.12, x + 3.5, 5.12,
                     MUTED, 1.35, arrow=True)

    add_line(slide, 17.0, 5.98, 17.0, 6.32, ACCENT, 1.5, arrow=True)
    add_card(slide, 13.5, 6.43, 5.1, 1.46,
             "Decision rule", "Score ≥ 0.45 on two consecutive evaluations\n1 s cooldown",
             accent=True, fill=PALE, title_size=16, body_size=13)
    add_text(slide, 0.96, 6.43, 11.8, 1.16,
             "No gain, AGC, denoising or noise conditioning is applied before model inference.\n"
             "AC-RMS is used for the upload endpoint only; it does not alter model PCM.",
             14, INK, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.96, 8.2, 17.5, 0.4,
             "This is the exported model's required feature path—not raw PCM inference and not a newly trained model in the firmware repository.",
             12, MUTED)
    return slide


def slide_5(prs):
    slide = new_slide(prs)
    add_header(slide, "Microphone capture and activation-gated PCM transport",
               "04  /  Hardware and audio path", 5)

    add_text(slide, 0.93, 1.58, 6.9, 0.3, "PROTOTYPE HARDWARE", 10.5, MUTED, True)
    hardware = [
        ("EDGE BOARD", "ESP32-S3 N16R8\n16 MB flash · 8 MB PSRAM"),
        ("MICROPHONE", "INMP441 I²S\nBCLK GPIO4 · WS GPIO5 · SD GPIO6 · L/R GND"),
        ("NETWORK", "2.4-GHz Wi-Fi\nWebSocket to backend host:8765 /v1/stream"),
        ("OPTIONAL SPEAKER", "MAX98357A I²S TX: BCLK4 / LRC5 / DIN8\nAcoustic result unverified; measured VIN was ~2.5 V"),
    ]
    y = 2.02
    for i, (label, body) in enumerate(hardware):
        add_rect(slide, 0.93, y, 7.35, 1.16,
                 PALE if i % 2 == 0 else WHITE, LINE, 0.9)
        add_text(slide, 1.15, y + 0.16, 1.72, 0.3, label,
                 10.5, ACCENT, True)
        add_text(slide, 2.98, y + 0.12, 5.05, 0.8, body,
                 13, INK, valign=MSO_ANCHOR.MIDDLE)
        y += 1.29

    add_text(slide, 9.0, 1.58, 9.7, 0.3,
             "ONE WEBSOCKET BINARY MESSAGE  ·  EVERY 20 ms", 10.5, MUTED, True)
    add_rect(slide, 9.0, 2.13, 2.2, 1.23, PALE_BLUE, LINE, 1.0)
    add_text(slide, 9.16, 2.39, 1.86, 0.36,
             "16-byte\nheader", 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 11.2, 2.13, 6.95, 1.23, WHITE, ACCENT, 1.4)
    add_text(slide, 11.48, 2.39, 6.4, 0.34,
             "640-byte PCM16 payload", 17, NAVY, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, 11.47, 2.78, 6.4, 0.3,
             "320 mono samples · 16 kHz · signed little-endian",
             12.5, SLATE, align=PP_ALIGN.CENTER)
    add_text(slide, 9.06, 3.58, 9.0, 0.38,
             "656 application bytes per audio message (before WebSocket / TCP / Wi-Fi overhead)",
             12, MUTED)

    add_text(slide, 9.0, 4.48, 8.5, 0.3, "SESSION LIFECYCLE", 10.5, MUTED, True)
    steps = [
        (9.02, "LISTEN", "KWS runs locally"),
        (11.31, "TRIGGER", "Hello Tors accepted"),
        (13.6, "STREAM", "800 ms prebuffer + live PCM"),
        (15.89, "STOP", "Ambient-noise endpoint"),
    ]
    for i, (x, title, body) in enumerate(steps):
        add_rect(slide, x, 4.98, 2.1, 1.35,
                 PALE_BLUE if i == 2 else PALE,
                 ACCENT if i == 2 else LINE, 1.0)
        add_text(slide, x + 0.12, 5.16, 1.86, 0.28,
                 title, 11, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.11, 5.52, 1.88, 0.62,
                 body, 11.5, SLATE, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            add_line(slide, x + 2.1, 5.63, x + 2.25, 5.63,
                     ACCENT if i >= 1 else MUTED, 1.3, arrow=True)

    add_rect(slide, 0.95, 8.0, 17.25, 1.0, PALE_BLUE, LINE, 0.9)
    add_text(slide, 1.2, 8.22, 16.7, 0.55,
             "Transport uses uncompressed PCM; the rolling prebuffer can include the wake phrase.\n"
             "The endpoint is energy-based, not a trained voice-activity detector.",
             13, INK, valign=MSO_ANCHOR.MIDDLE)
    return slide


def slide_6(prs):
    slide = new_slide(prs)
    add_header(slide, "Prototype status: integrated paths, open evaluation", 
               "05  /  Feasibility and status", 6)
    columns = [
        (0.93, 5.63, "IMPLEMENTED IN FIRMWARE", PALE_BLUE),
        (6.84, 5.63, "HOST / BUILD EVIDENCE", PALE),
        (12.75, 6.0, "NOT YET ESTABLISHED", PALE),
    ]
    for x, w, label, fill in columns:
        add_rect(slide, x, 1.74, w, 0.62, fill, LINE, 0.9)
        add_text(slide, x + 0.14, 1.9, w - 0.28, 0.3,
                 label, 12, NAVY, True, align=PP_ALIGN.CENTER)
    implemented = (
        "• Custom Hello Tors KWS on ESP32-S3\n"
        "• INMP441 capture and PCM conversion\n"
        "• Trigger-gated WebSocket streaming\n"
        "• Device RAM / MIC / KWS / HEALTH logs\n"
        "• Local KWS-triggered WAV recordings\n"
        "• MAX98357A firmware transmit path"
    )
    host = (
        "• Model artifact identity / integration checks\n"
        "• Host feature vectors: 1,960 / 1,960 exact\n"
        "• Streaming frontend parity checks\n"
        "• WebSocket framing and local backend tests\n"
        "• Boot self-test PASS observed in device log\n"
        "• Latest documented firmware builds"
    )
    pending = (
        "• Labeled INMP441 detection-rate trials\n"
        "• False activations per hour under noise\n"
        "• Whole-system RAM on the latest flashed build\n"
        "• Idle CPU utilization measurement\n"
        "• Keyword-end → server-receipt p50/p95/p99\n"
        "• Verified acoustic speaker output"
    )
    for (x, w, _, _), content in zip(columns, (implemented, host, pending)):
        add_rect(slide, x, 2.48, w, 4.55, WHITE, LINE, 0.9)
        add_text(slide, x + 0.2, 2.73, w - 0.38, 4.03,
                 content, 14, INK, line_spacing=1.13)
    add_rect(slide, 0.95, 7.48, 17.82, 1.16, PALE_BLUE, LINE, 0.9)
    add_text(slide, 1.2, 7.71, 2.15, 0.3, "DESIGN INTENT", 11, ACCENT, True)
    add_text(slide, 3.34, 7.62, 15.0, 0.72,
             "Keep the wake decision on-device; send speech only after a local trigger."
             " The audio path is implemented, but accuracy, resource-limit compliance and end-to-end latency require controlled trials.",
             14, INK, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.98, 9.05, 17.4, 0.42,
             "Backend distinction: the separate server source has a Vosk worker; local_backend.py is an ASR-free playback/transport test tool.",
             11.5, MUTED)
    return slide


def slide_7(prs):
    slide = new_slide(prs)
    add_header(slide, "Current linked memory footprint — static allocations only",
               "06  /  Technical supplement: resources", 7)
    add_text(slide, 0.95, 1.57, 17.6, 0.48,
             "Latest documented build · ESP32-S3 N16R8 · 16 MB flash / 8 MB PSRAM",
             15, SLATE)

    x0, y0 = 0.95, 2.25
    widths = [6.0, 2.5, 9.3]
    headers = ["ITEM", "BYTES", "PLACEMENT / INTERPRETATION"]
    x = x0
    for width, label in zip(widths, headers):
        add_rect(slide, x, y0, width, 0.62, NAVY, NAVY, 0)
        add_text(slide, x + 0.12, y0 + 0.15, width - 0.24, 0.3,
                 label, 11, WHITE, True,
                 align=PP_ALIGN.RIGHT if label == "BYTES" else PP_ALIGN.LEFT)
        x += width
    rows = [
        ("Embedded INT8 model", "36,744", "Mapped from flash; model weights only"),
        ("Application binary", "1,227,280", "Flash image; includes runtime, diagnostics, networking and other code/data"),
        ("Tensor arena reserved", "114,688", "112 KiB in internal BSS; already included below"),
        ("Internal static data + BSS", "163,668", "Includes tensor arena, plus other linked internal static data"),
        ("External static BSS", "96,064", "PSRAM; excludes dynamic PSRAM heap use"),
        ("Static BSS combined", "259,732", "Internal + external static only; excludes runtime heap, stacks and allocator overhead"),
    ]
    y = y0 + 0.62
    row_h = 0.79
    for i, row in enumerate(rows):
        fill = PALE_BLUE if i == len(rows) - 1 else (WHITE if i % 2 else PALE)
        x = x0
        for j, (width, value) in enumerate(zip(widths, row)):
            add_rect(slide, x, y, width, row_h, fill, LINE, 0.65)
            add_text(slide, x + 0.12, y + 0.13, width - 0.24, row_h - 0.18,
                     value, 13.5 if j != 1 else 14, NAVY if j == 1 else INK,
                     j == 1 or i == len(rows) - 1,
                     align=PP_ALIGN.RIGHT if j == 1 else PP_ALIGN.LEFT,
                     valign=MSO_ANCHOR.MIDDLE)
            x += width
        y += row_h

    add_rect(slide, 0.97, 8.12, 17.8, 1.12, PALE_BLUE, ACCENT, 1.1)
    add_text(slide, 1.23, 8.29, 3.6, 0.38,
             "<256 KB RAM TARGET", 14, ACCENT, True)
    add_text(slide, 4.85, 8.23, 13.4, 0.68,
             "Full-runtime compliance is not demonstrated. Static allocations alone are 259,732 B across internal SRAM and PSRAM;"
             " dynamic heaps and task stacks are additional. Do not add the 114,688-byte arena twice.",
             12.5, INK, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.98, 9.58, 17.6, 0.35,
             "Older serial heap/arena snapshots belong to a different linked footprint and are intentionally excluded from this current-build table.",
             10.5, MUTED)
    return slide


def slide_8(prs):
    slide = new_slide(prs)
    add_header(slide, "Evaluation evidence and next measurements",
               "07  /  Technical supplement: validation", 8)
    x0, y0 = 0.93, 1.82
    widths = [5.1, 4.45, 8.45]
    labels = ["METRIC", "CURRENT EVIDENCE", "NEXT MEASUREMENT"]
    x = x0
    for width, label in zip(widths, labels):
        add_rect(slide, x, y0, width, 0.58, NAVY, NAVY, 0)
        add_text(slide, x + 0.14, y0 + 0.14, width - 0.28, 0.28,
                 label, 10.5, WHITE, True)
        x += width
    rows = [
        ("Model footprint", "36,744-byte model; 114,688-byte arena reserved", "Report flash and current arena used on the flashed build"),
        ("Feature correctness", "Host goldens: 1,960 / 1,960 exact per vector", "Confirm boot self-tests and compare captured mic WAV against ML reference"),
        ("Keyword accuracy", "Threshold 0.45; two consecutive evaluations", "Labeled positive trials across speakers, distance and noise"),
        ("False activations", "No valid per-hour field result", "Hours of labeled silence, fan, music and unrelated speech"),
        ("Idle CPU", "<10% is a requirement; not measured", "Measure total continuous-listening CPU, not invoke time alone"),
        ("End-to-end latency", "Backend reports start→first-frame proxy only", "Measure keyword end→server receipt; collect p50 / p95 / p99"),
        ("Speaker output", "I²S transmit path in firmware; acoustics unverified", "Resolve MAX98357A supply and verify playback at the speaker"),
    ]
    y = y0 + 0.58
    row_h = 0.78
    for i, row in enumerate(rows):
        fill = WHITE if i % 2 else PALE
        x = x0
        for j, (width, value) in enumerate(zip(widths, row)):
            add_rect(slide, x, y, width, row_h, fill, LINE, 0.6)
            add_text(slide, x + 0.12, y + 0.09, width - 0.24, row_h - 0.14,
                     value, 11.5 if j else 12.3,
                     NAVY if j == 0 else INK, j == 0,
                     valign=MSO_ANCHOR.MIDDLE)
            x += width
        y += row_h

    add_text(slide, 0.97, 8.25, 17.7, 0.26,
             "REFERENCES  /  IMPLEMENTATION SOURCES", 10, ACCENT, True)
    add_text(slide, 0.97, 8.58, 17.7, 1.03,
             "SIH 26172 problem statement (ISRO)  ·  docs/ESP32S3_ML_DEPLOYMENT_REVIEW.md  ·  "
             "docs/ESP32_BACKEND_DEVICE_CONTRACT.md  ·  docs/LOCAL_DIAGNOSTICS.md  ·  "
             "docs/SPEAKER_PLAYBACK_FIX.md  ·  SERVER SIDE CODE/sih-voice-activator-main/{ARCHITECTURE,PROTOCOL,README}.md",
             10.3, SLATE, line_spacing=1.1)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH_IN)
    prs.slide_height = Inches(HEIGHT_IN)
    prs.core_properties.title = "EdgeAIKWS — SIH 26172"
    prs.core_properties.subject = "Low Latency and Efficient Voice Activator for Edge Devices"
    prs.core_properties.author = "EdgeAIKWS project"
    prs.core_properties.keywords = "SIH 26172, ESP32-S3, KWS, Hello Tors, TFLite Micro"

    slide_1(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)
    slide_6(prs)
    slide_7(prs)
    slide_8(prs)
    prs.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
