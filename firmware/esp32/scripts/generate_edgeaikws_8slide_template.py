"""Generate an editable, placeholder-based eight-slide SIH working deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "EdgeAIKWS_SIH26172_8Slide_Template.pptx"

WIDTH = 20.0
HEIGHT = 11.25
FONT = "Arial"

NAVY = "20364F"
SLATE = "566777"
INK = "263541"
MUTED = "73818D"
PALE = "F2F5F7"
PALE_BLUE = "E8EFF4"
LINE = "CCD5DC"
ORANGE = "B96B43"
WHITE = "FFFFFF"


def color(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def add_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    value: str,
    size: float = 15,
    fill: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.06,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(max(2, size * 0.18))
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
    return shape


def add_rect(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    fill: str = WHITE,
    line: str = LINE,
    line_width: float = 1.0,
    shape_type=MSO_SHAPE.RECTANGLE,
):
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line)
    shape.line.width = Pt(line_width)
    return shape


def add_rule(slide, x1, y1, x2, y2, fill=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    return shape


def new_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(WHITE)
    return slide


def add_header(slide, heading: str, number: int):
    add_text(slide, 0.78, 0.28, 12.0, 0.25, "EDITABLE SIH WORKING TEMPLATE", 9, ORANGE, True)
    add_text(slide, 0.76, 0.57, 18.5, 0.64, heading, 27, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 0.8, 1.34, 18.4, 0.025, LINE, LINE, 0)
    add_footer(slide, number)


def add_footer(slide, number: int):
    add_text(
        slide,
        0.82,
        10.76,
        15.2,
        0.32,
        "WORKING TEMPLATE  ·  REPLACE BRACKETED FIELDS",
        8.5,
        MUTED,
        True,
        margin=0.04,
    )
    add_text(
        slide,
        18.1,
        10.76,
        1.05,
        0.32,
        f"{number:02d} / 08",
        9,
        MUTED,
        align=PP_ALIGN.RIGHT,
        margin=0.04,
    )


def add_panel(slide, x, y, width, height, title, fill=PALE, title_size=13):
    add_rect(slide, x, y, width, height, fill, LINE, 0.9)
    add_text(slide, x + 0.16, y + 0.12, width - 0.32, 0.34, title, title_size, NAVY, True)


def add_card(slide, x, y, width, height, title, body, accent=False, fill=WHITE):
    add_rect(slide, x, y, width, height, fill, LINE, 0.9)
    if accent:
        add_rect(slide, x, y, 0.07, height, ORANGE, ORANGE, 0)
    inset = 0.2 if accent else 0.15
    add_text(slide, x + inset, y + 0.14, width - inset - 0.12, 0.38, title, 14, NAVY, True)
    add_text(slide, x + inset, y + 0.62, width - inset - 0.14, height - 0.74, body, 12, SLATE)


def add_native_table(slide, x, y, width, height, rows, column_widths, font_size=11.5):
    table_shape = slide.shapes.add_table(
        len(rows), len(column_widths), Inches(x), Inches(y), Inches(width), Inches(height)
    )
    table = table_shape.table
    for index, column_width in enumerate(column_widths):
        table.columns[index].width = Inches(column_width)
    row_height = height / len(rows)
    for row_index, row_data in enumerate(rows):
        table.rows[row_index].height = Inches(row_height)
        for column_index, value in enumerate(row_data):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.margin_left = Inches(0.11)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(NAVY if row_index == 0 else (WHITE if row_index % 2 else PALE))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.font.name = FONT
                paragraph.font.size = Pt(font_size if row_index else font_size - 0.5)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = color(WHITE if row_index == 0 else INK)
    return table_shape


def add_image_frame(slide, x, y, width, height, prompt):
    frame = add_rect(slide, x, y, width, height, PALE, MUTED, 1.2)
    frame.line.dash_style = 2
    add_text(
        slide,
        x + 0.25,
        y + height / 2 - 0.32,
        width - 0.5,
        0.66,
        prompt,
        15,
        SLATE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return frame


def slide_1(presentation):
    slide = new_slide(presentation)
    add_text(slide, 0.82, 0.28, 2.5, 0.25, "TITLE PAGE", 10, ORANGE, True)
    add_rect(slide, 0.82, 0.68, 2.4, 1.05, PALE_BLUE, LINE)
    add_text(slide, 0.95, 0.94, 2.14, 0.48, "[Add team identity]", 13, SLATE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 3.65, 0.72, 11.0, 0.6, "SMART INDIA HACKATHON", 25, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 6.1, 1.3, 6.1, 0.28, "[Add edition / year]", 11, MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, 16.72, 0.68, 2.4, 1.05, PALE_BLUE, LINE)
    add_text(slide, 16.87, 0.94, 2.1, 0.48, "[Add SIH identity]", 12, SLATE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

    add_text(slide, 0.88, 2.2, 9.3, 0.36, "PROJECT METADATA", 11, ORANGE, True)
    metadata = [
        ("Problem statement ID", "[Add problem statement ID]"),
        ("Organization", "[Add organization]"),
        ("Theme / category", "[Add theme and category]"),
        ("Team", "[Add team name and lead]"),
    ]
    for index, (label, value) in enumerate(metadata):
        y = 2.78 + index * 0.83
        add_text(slide, 0.9, y, 3.1, 0.22, label.upper(), 9, MUTED, True)
        add_text(slide, 0.9, y + 0.24, 8.7, 0.42, value, 15, INK, True)

    add_image_frame(slide, 10.45, 2.1, 8.65, 6.1, "[Add project image or mark]")
    add_rect(slide, 0.88, 7.45, 0.9, 0.9, PALE_BLUE, LINE)
    add_text(slide, 0.96, 7.67, 0.74, 0.42, "[Mark]", 10, SLATE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 2.02, 7.48, 7.85, 0.92, "[Add project title]", 26, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    add_footer(slide, 1)
    return slide


def slide_2(presentation):
    slide = new_slide(presentation)
    add_header(slide, "IDEA TITLE", 2)
    add_rect(slide, 0.82, 1.62, 18.36, 1.2, PALE_BLUE, LINE)
    add_text(slide, 1.02, 1.76, 4.1, 0.25, "IDEA TITLE", 10, ORANGE, True)
    add_text(slide, 1.02, 2.06, 4.1, 0.5, "[Add concise idea title]", 17, NAVY, True)
    add_rule(slide, 5.45, 1.82, 5.45, 2.62, LINE, 1.0)
    add_text(slide, 5.75, 1.76, 3.4, 0.25, "PROPOSED SOLUTION", 10, ORANGE, True)
    add_text(slide, 5.75, 2.05, 12.95, 0.58, "[Describe the proposed solution in one or two editable sentences]", 15, INK)

    add_text(slide, 0.86, 3.08, 8.0, 0.3, "FOUR STEP PROCESS", 10, MUTED, True)
    step_y, step_w, gap = 3.52, 4.35, 0.27
    for index in range(4):
        x = 0.82 + index * (step_w + gap)
        add_rect(slide, x, step_y, step_w, 1.75, PALE if index % 2 == 0 else WHITE, LINE)
        add_text(slide, x + 0.16, step_y + 0.14, 0.55, 0.3, f"0{index + 1}", 11, ORANGE, True)
        add_text(slide, x + 0.77, step_y + 0.12, step_w - 0.92, 0.42, f"[Add step {index + 1} title]", 14, NAVY, True)
        add_text(slide, x + 0.18, step_y + 0.68, step_w - 0.36, 0.88, "[Add a short description of this process step]", 11.5, SLATE)

    add_text(slide, 0.86, 5.63, 9.0, 0.3, "SUMMARY / COMPARISON", 10, MUTED, True)
    rows = [
        ["DIMENSION", "[Add existing approach]", "[Add proposed approach]"],
        ["[Add comparison point]", "[Add summary]", "[Add summary]"],
        ["[Add comparison point]", "[Add summary]", "[Add summary]"],
        ["[Add comparison point]", "[Add summary]", "[Add summary]"],
    ]
    add_native_table(slide, 0.82, 6.02, 18.36, 4.35, rows, [4.6, 6.88, 6.88], 12)
    return slide


def slide_3(presentation):
    slide = new_slide(presentation)
    add_header(slide, "TECHNICAL APPROACH", 3)
    add_panel(slide, 0.82, 1.62, 13.65, 6.72, "ARCHITECTURE / PROCESS FLOW", PALE)
    nodes = [
        "[Add input or\nsensor]",
        "[Add processing\nstage]",
        "[Add model or\ndecision stage]",
        "[Add output or\ninterface]",
        "[Add system\noutcome]",
    ]
    start_x, node_w, gap, node_y, node_h = 1.08, 2.25, 0.38, 4.03, 1.55
    for index, label in enumerate(nodes):
        x = start_x + index * (node_w + gap)
        add_rect(slide, x, node_y, node_w, node_h, WHITE, ORANGE if index == 2 else LINE, 1.15)
        add_text(slide, x + 0.13, node_y + 0.2, node_w - 0.26, 1.1, label, 13, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if index < len(nodes) - 1:
            line_start = x + node_w
            line_end = line_start + gap - 0.1
            add_rule(slide, line_start + 0.02, node_y + node_h / 2, line_end, node_y + node_h / 2, ORANGE, 1.5)
            arrow = add_rect(slide, line_end - 0.05, node_y + node_h / 2 - 0.07, 0.14, 0.14, ORANGE, ORANGE, 0, MSO_SHAPE.ISOSCELES_TRIANGLE)
            arrow.rotation = 90
    add_text(slide, 1.12, 2.34, 12.9, 0.65, "[Add a short explanation of data flow and component interactions]", 13, SLATE)

    add_panel(slide, 14.78, 1.62, 4.4, 6.72, "TECH STACK", PALE_BLUE)
    stack = ["[Add hardware]", "[Add framework / runtime]", "[Add model or algorithm]", "[Add software interface]", "[Add language / tools]"]
    for index, item in enumerate(stack):
        y = 2.25 + index * 1.08
        add_text(slide, 14.98, y, 3.98, 0.23, f"COMPONENT {index + 1}", 8.5, MUTED, True)
        add_text(slide, 14.98, y + 0.25, 3.98, 0.53, item, 12, NAVY, True)
        if index < len(stack) - 1:
            add_rule(slide, 14.98, y + 0.9, 18.98, y + 0.9, LINE, 0.8)

    add_rect(slide, 0.82, 8.62, 18.36, 1.75, PALE, LINE)
    add_text(slide, 1.02, 8.79, 2.5, 0.28, "NOTES / CALLOUT", 10, ORANGE, True)
    add_text(slide, 3.22, 8.77, 15.6, 1.3, "[Add assumptions, interfaces, constraints, or technical notes]", 13, SLATE, valign=MSO_ANCHOR.MIDDLE)
    return slide


def slide_4(presentation):
    slide = new_slide(presentation)
    add_header(slide, "ML MODEL PARAMETERS & ESP32 DEPLOYMENT PARAMETERS", 4)
    add_panel(slide, 0.82, 1.65, 9.0, 8.7, "ML MODEL PARAMETERS", PALE_BLUE)
    add_panel(slide, 10.18, 1.65, 9.0, 8.7, "ESP32 DEPLOYMENT PARAMETERS", PALE)
    left_rows = [
        ["PARAMETER", "VALUE / NOTES"],
        ["[Add model name / type]", "[Add value]"],
        ["[Add input shape / format]", "[Add value]"],
        ["[Add training data / source]", "[Add value]"],
        ["[Add training configuration]", "[Add value]"],
        ["[Add evaluation metric]", "[Add value / evidence]"],
        ["[Add model size / quantization]", "[Add value]"],
    ]
    right_rows = [
        ["PARAMETER", "VALUE / NOTES"],
        ["[Add ESP32 variant]", "[Add value]"],
        ["[Add clock / memory settings]", "[Add value]"],
        ["[Add inference runtime]", "[Add value]"],
        ["[Add input / output interface]", "[Add value]"],
        ["[Add latency / resource use]", "[Add value / evidence]"],
        ["[Add power / deployment setting]", "[Add value]"],
    ]
    add_native_table(slide, 1.02, 2.27, 8.6, 7.82, left_rows, [4.0, 4.6], 11.5)
    add_native_table(slide, 10.38, 2.27, 8.6, 7.82, right_rows, [4.0, 4.6], 11.5)
    return slide


def slide_5(presentation):
    slide = new_slide(presentation)
    add_header(slide, "DASHBOARD IMAGES & HARDWARE IMAGES", 5)
    add_text(slide, 0.86, 1.57, 10.7, 0.25, "DASHBOARD", 10, MUTED, True)
    add_image_frame(slide, 0.82, 1.9, 10.85, 7.45, "[Add dashboard screenshot]")
    add_text(slide, 0.9, 9.53, 10.65, 0.62, "[Add dashboard image caption / context]", 12, SLATE)

    add_text(slide, 12.05, 1.57, 7.05, 0.25, "HARDWARE", 10, MUTED, True)
    add_image_frame(slide, 12.02, 1.9, 7.16, 3.25, "[Add hardware image 1]")
    add_text(slide, 12.1, 5.21, 7.0, 0.47, "[Add hardware image 1 caption]", 11, SLATE)
    add_image_frame(slide, 12.02, 5.84, 7.16, 3.25, "[Add hardware image 2]")
    add_text(slide, 12.1, 9.16, 7.0, 0.47, "[Add hardware image 2 caption]", 11, SLATE)
    return slide


def slide_6(presentation):
    slide = new_slide(presentation)
    add_header(slide, "FEASIBILITY & VIABILITY", 6)
    add_text(slide, 0.86, 1.56, 10.0, 0.25, "COMPARISON WITH EXISTING PRODUCTS / APPROACHES", 10, MUTED, True)
    comparison = [
        ["DIMENSION", "[Add existing option]", "[Add proposed solution]", "[Add evidence / source]"],
        ["[Add dimension]", "[Add summary]", "[Add summary]", "[Add source]"],
        ["[Add dimension]", "[Add summary]", "[Add summary]", "[Add source]"],
    ]
    add_native_table(slide, 0.82, 1.9, 18.36, 2.45, comparison, [3.55, 4.75, 5.0, 5.06], 10.5)

    panel_y, panel_h, panel_w, panel_gap = 4.7, 5.65, 5.9, 0.33
    panel_xs = [0.82, 0.82 + panel_w + panel_gap, 0.82 + 2 * (panel_w + panel_gap)]
    titles = ["FEASIBILITY", "VIABILITY", "CHALLENGES / MITIGATIONS"]
    for x, title in zip(panel_xs, titles):
        add_panel(slide, x, panel_y, panel_w, panel_h, title, PALE_BLUE if title == "VIABILITY" else PALE)
    add_text(slide, panel_xs[0] + 0.18, 5.36, panel_w - 0.36, 4.65,
             "TECHNICAL\n[Add technical evidence]\n\nOPERATIONAL\n[Add operational requirements]\n\nRESOURCES\n[Add required resources]", 12, SLATE)
    add_text(slide, panel_xs[1] + 0.18, 5.36, panel_w - 0.36, 4.65,
             "USERS / CONTEXT\n[Add intended use context]\n\nCOST / SUSTAINABILITY\n[Add evidence or estimate]\n\nADOPTION\n[Add deployment considerations]", 12, SLATE)
    add_text(slide, panel_xs[2] + 0.18, 5.36, panel_w - 0.36, 4.65,
             "CHALLENGE\n[Add challenge]  →  [Add mitigation]\n\nCHALLENGE\n[Add challenge]  →  [Add mitigation]\n\nCHALLENGE\n[Add challenge]  →  [Add mitigation]", 11.5, SLATE)
    return slide


def slide_7(presentation):
    slide = new_slide(presentation)
    add_header(slide, "IMPACT & BENEFITS", 7)
    card_w, card_h, gap_x, gap_y = 5.94, 2.43, 0.28, 0.25
    titles = ["[Add benefit 1]", "[Add benefit 2]", "[Add benefit 3]", "[Add benefit 4]", "[Add benefit 5]", "[Add benefit 6]"]
    for index, title in enumerate(titles):
        row, column = divmod(index, 3)
        x = 0.82 + column * (card_w + gap_x)
        y = 1.68 + row * (card_h + gap_y)
        add_card(slide, x, y, card_w, card_h, title, "[Add expected benefit, impact, and supporting evidence]", accent=(index == 0), fill=PALE if index % 2 == 0 else WHITE)

    add_rect(slide, 0.82, 6.94, 18.36, 1.47, PALE_BLUE, LINE)
    add_text(slide, 1.02, 7.09, 2.65, 0.3, "BENEFICIARIES", 10, ORANGE, True)
    for index in range(4):
        x = 3.85 + index * 3.72
        add_text(slide, x, 7.09, 3.42, 0.28, f"[Add beneficiary group {index + 1}]", 11.5, NAVY, True)
        add_text(slide, x, 7.44, 3.42, 0.65, "[Add relevance / value]", 10.5, SLATE)

    add_rect(slide, 0.82, 8.72, 18.36, 1.56, PALE, LINE)
    add_text(slide, 1.02, 8.88, 2.65, 0.3, "OPTIONAL NEXT STEP", 10, ORANGE, True)
    add_text(slide, 3.85, 8.84, 14.95, 1.1, "[Add a proposed next step, owner, or validation milestone]", 13, SLATE, valign=MSO_ANCHOR.MIDDLE)
    return slide


def slide_8(presentation):
    slide = new_slide(presentation)
    add_header(slide, "RESEARCH & REFERENCES", 8)
    add_panel(slide, 0.82, 1.62, 11.55, 7.75, "REFERENCE LIST", PALE)
    for index in range(6):
        y = 2.2 + index * 1.05
        add_text(slide, 1.05, y, 0.52, 0.32, f"[{index + 1}]", 11, ORANGE, True)
        add_text(slide, 1.62, y, 10.35, 0.8, "[Add citation: author, title, venue, year, and persistent link]", 12, INK)
        if index < 5:
            add_rule(slide, 1.05, y + 0.88, 12.1, y + 0.88, LINE, 0.7)

    add_panel(slide, 12.72, 1.62, 6.46, 7.75, "FOUNDATIONAL WORK / SOURCES", PALE_BLUE)
    source_fields = [
        ("FOUNDATIONAL WORK", "[Add prior work or method]"),
        ("DATA / MATERIALS", "[Add dataset or material source]"),
        ("TOOLS / FRAMEWORKS", "[Add tool and version]"),
        ("ATTRIBUTION / LICENSE", "[Add attribution and license details]"),
    ]
    for index, (label, value) in enumerate(source_fields):
        y = 2.28 + index * 1.48
        add_text(slide, 12.96, y, 5.95, 0.24, label, 9, MUTED, True)
        add_text(slide, 12.96, y + 0.32, 5.9, 0.75, value, 12, NAVY, True)
        if index < len(source_fields) - 1:
            add_rule(slide, 12.96, y + 1.2, 18.9, y + 1.2, LINE, 0.8)

    add_text(slide, 0.86, 9.63, 2.5, 0.25, "EDITABLE LINKS", 9, ORANGE, True)
    link_prompts = ["[Add source URL]", "[Add source URL]", "[Add source URL]"]
    for index, prompt in enumerate(link_prompts):
        x = 3.0 + index * 5.37
        add_rect(slide, x, 9.56, 5.05, 0.58, WHITE, LINE, 0.9)
        add_text(slide, x + 0.12, 9.68, 4.8, 0.3, prompt, 10.5, SLATE)
    return slide


def build_presentation():
    presentation = Presentation()
    presentation.slide_width = Inches(WIDTH)
    presentation.slide_height = Inches(HEIGHT)
    for make_slide in (
        slide_1,
        slide_2,
        slide_3,
        slide_4,
        slide_5,
        slide_6,
        slide_7,
        slide_8,
    ):
        make_slide(presentation)
    return presentation


def main():
    presentation = build_presentation()
    presentation.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
