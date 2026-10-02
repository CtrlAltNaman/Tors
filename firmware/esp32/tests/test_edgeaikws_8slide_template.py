import importlib.util
from pathlib import Path
import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "EdgeAIKWS_SIH26172_8Slide_Template.pptx"
EXPECTED_TABLE_SLIDES = (2, 4, 6)
EXPECTED_HEADINGS = [
    "TITLE PAGE",
    "IDEA TITLE",
    "TECHNICAL APPROACH",
    "ML MODEL PARAMETERS & ESP32 DEPLOYMENT PARAMETERS",
    "DASHBOARD IMAGES & HARDWARE IMAGES",
    "FEASIBILITY & VIABILITY",
    "IMPACT & BENEFITS",
    "RESEARCH & REFERENCES",
]


def slide_text(slide: object) -> str:
    text_parts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text_parts.append(shape.text)
        if getattr(shape, "has_table", False):
            text_parts.extend(
                cell.text for row in shape.table.rows for cell in row.cells
            )
    return "\n".join(text_parts)


class DeckAcceptanceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.assertTrue(DECK.is_file(), f"Generated deck is missing: {DECK}")
        self.presentation = Presentation(DECK)

    def test_deck_has_eight_slides_in_requested_order(self) -> None:
        self.assertEqual(len(self.presentation.slides), 8)
        for slide, heading in zip(
            self.presentation.slides, EXPECTED_HEADINGS, strict=True
        ):
            self.assertIn(heading, slide_text(slide))

    def test_every_slide_has_an_editable_add_placeholder(self) -> None:
        for slide_number, slide in enumerate(self.presentation.slides, start=1):
            self.assertRegex(
                slide_text(slide),
                r"\[Add [^\]]+\]",
                f"Slide {slide_number} has no editable [Add ...] placeholder",
            )

    def test_deck_uses_wide_canvas_and_editable_content(self) -> None:
        self.assertEqual(round(self.presentation.slide_width / 914400, 2), 20.00)
        self.assertEqual(round(self.presentation.slide_height / 914400, 2), 11.25)
        for slide_number, slide in enumerate(self.presentation.slides, start=1):
            self.assertGreaterEqual(len(slide.shapes), 6)
            self.assertTrue(any(shape.has_text_frame for shape in slide.shapes))
            self.assertFalse(
                any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes),
                f"Slide {slide_number} contains a flattened picture shape",
            )
            if slide_number in EXPECTED_TABLE_SLIDES:
                self.assertTrue(
                    any(shape.has_table for shape in slide.shapes),
                    f"Slide {slide_number} should contain a native PowerPoint table",
                )

    def test_deck_contains_no_dhvani_specific_branding_or_claims(self) -> None:
        all_text = "\n".join(
            slide_text(slide) for slide in self.presentation.slides
        ).casefold()

        for forbidden in ("dhvani", "money.exe", "dccrn", "noise cancellation"):
            self.assertNotIn(forbidden, all_text)

    def test_slide_text_includes_native_table_cell_text(self) -> None:
        scratch = Presentation()
        slide = scratch.slides.add_slide(scratch.slide_layouts[6])
        table = slide.shapes.add_table(1, 1, 0, 0, 914400, 914400).table
        table.cell(0, 0).text = "TABLE_ONLY_CONTENT"

        self.assertIn("TABLE_ONLY_CONTENT", slide_text(slide))

    def test_footer_and_page_number_boxes_have_readable_height(self) -> None:
        minimum_height = 0.30 * 914400
        for slide in self.presentation.slides:
            footer = next(
                shape
                for shape in slide.shapes
                if shape.has_text_frame and "REPLACE BRACKETED FIELDS" in shape.text
            )
            page_number = next(
                shape
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip().endswith("/ 08")
            )
            self.assertGreaterEqual(footer.height, minimum_height)
            self.assertGreaterEqual(page_number.height, minimum_height)

    def test_existing_reference_and_content_decks_are_preserved(self) -> None:
        dhvani_deck = ROOT / "Dhvani_Kavach_SIH26052.pptx"
        content_deck = ROOT / "EdgeAIKWS_SIH26172.pptx"
        self.assertTrue(dhvani_deck.is_file())
        self.assertTrue(content_deck.is_file())

        generator_path = ROOT / "scripts" / "generate_edgeaikws_8slide_template.py"
        spec = importlib.util.spec_from_file_location(
            "edgeaikws_8slide_template_generator", generator_path
        )
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        generator = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(generator)

        output = Path(generator.OUTPUT).resolve()
        self.assertEqual(output, DECK.resolve())
        self.assertNotEqual(output, dhvani_deck.resolve())
        self.assertNotEqual(output, content_deck.resolve())


if __name__ == "__main__":
    unittest.main()
